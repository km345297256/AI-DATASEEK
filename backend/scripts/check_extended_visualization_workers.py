"""Exercise real v2 Docker readers without files, sessions, jobs or model calls.

Run from an existing one-off backend container with its configured Docker socket
and ``PYTHONPATH`` pointing to this backend checkout. The sandbox image must be
built first. No service/port is created and no user file or database is accessed.

A separate, constrained fixture producer uses installed scientific libraries to
generate small synthetic files. Every preview then goes through the *real*
backend Docker stdin/framing gateway and production ``validate_payload``. The
client observer only inspects real containers; it never replaces a parser or a
Docker operation with a fake result. Owned containers are removed in ``finally``.
"""
from __future__ import annotations

import argparse
import base64
import contextlib
import io
import json
import multiprocessing
import threading
import time
import uuid
import zipfile
from typing import Any

import docker

from app.application.services.extended_visualization import MAX_OUTPUT, validate_options, validate_payload
from app.core.config import get_settings
from app.infrastructure.external.sandbox import extended_visualization_worker as gateway
from app.infrastructure.external.sandbox.visualization_worker import VisualizationWorkerError


FIXTURE_CODE = r'''
import base64, io, json, os, tempfile
from pathlib import Path
import numpy as np
import h5py
import openpyxl
import uproot
from scipy.io import savemat
from netCDF4 import Dataset
from docx import Document
from pptx import Presentation

assert os.getuid() == 65534
files = {}
def add(name, value):
    assert 0 < len(value) < 1024 * 1024
    files[name] = base64.b64encode(value).decode('ascii')

with tempfile.TemporaryDirectory(prefix='visualization-fixtures-') as directory:
    path = Path(directory)
    stream = io.BytesIO()
    np.save(stream, np.arange(24, dtype=float).reshape(2, 3, 4))
    add('array.npy', stream.getvalue())
    stream = io.BytesIO()
    np.savez(stream, temperature=np.arange(6).reshape(2, 3))
    add('array.npz', stream.getvalue())
    stream = io.BytesIO()
    savemat(stream, {'temperature': np.arange(6).reshape(2, 3)})
    add('array.mat', stream.getvalue())
    with h5py.File(path / 'sample.h5', 'w') as root:
        values = root.create_dataset('entry/signal', data=np.arange(24).reshape(2, 3, 4))
        values.attrs['units'] = np.bytes_('counts')
        root['outside'] = h5py.ExternalLink('/forbidden/external.h5', '/data')
    add('sample.h5', (path / 'sample.h5').read_bytes())
    for fmt in ['NETCDF4', 'NETCDF3_CLASSIC']:
        filename = path / (fmt + '.nc')
        with Dataset(filename, 'w', format=fmt) as root:
            root.createDimension('sample', 3)
            root.createVariable('signal', 'f4', ('sample',))[:] = [1.5, 2.5, 3.5]
        add(fmt + '.nc', filename.read_bytes())
    book = openpyxl.Workbook()
    book.active.title = 'Measurements'
    book.active.append(['time', 'temperature', 'formula'])
    book.active.append([1, 20.5, '=A2+B2'])
    stream = io.BytesIO()
    book.save(stream)
    add('book.xlsx', stream.getvalue())
    with uproot.recreate(path / 'histograms.root') as root:
        root['detector/counts'] = (np.array([1., 4., 9.]), np.array([0., 1., 3., 7.]))
        root['density'] = (np.arange(6, dtype=float).reshape(2, 3), np.array([0., 1., 4.]), np.array([-2., 0., 3., 8.]))
        root['curve'] = uproot.as_TGraph({'x': np.array([1., 3., 4.]), 'y': np.array([2., 7., 5.])})
        root['description'] = 'No TObjString, function or canvas execution'
    add('histograms.root', (path / 'histograms.root').read_bytes())
    document = Document()
    document.add_paragraph('DataSeek synthetic, isolated Word preview')
    stream = io.BytesIO()
    document.save(stream)
    add('document.docx', stream.getvalue())
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = 'DataSeek synthetic, isolated slide'
    stream = io.BytesIO()
    presentation.save(stream)
    add('slides.pptx', stream.getvalue())
print(json.dumps({'uid': os.getuid(), 'files': files}, separators=(',', ':')))
'''


JCAMP = b"""##TITLE=DataSeek synthetic processed spectrum
##JCAMP-DX=5.00
##DATA TYPE=NMR SPECTRUM
##.OBSERVE NUCLEUS=1H
##.OBSERVE FREQUENCY=400
##XUNITS=PPM
##NPOINTS=4
##FIRSTX=4
##LASTX=1
##YFACTOR=0.5
##XYDATA=(X++(Y..Y))
4 0 4
2 8 0
##END=
"""
METPY_OPTIONS = {"pressure_column": "p", "temperature_column": "t", "dewpoint_column": "td",
                 "pressure_unit": "hPa", "temperature_unit": "degC", "dewpoint_unit": "degC"}


def _inspect_isolation(container: Any) -> None:
    container.reload()
    attrs = container.attrs
    config, host = attrs["Config"], attrs["HostConfig"]
    assert config["User"] == "65534:65534", "Reader is not using the configured unprivileged UID"
    assert host["ReadonlyRootfs"] is True, "Reader root filesystem must be read-only"
    assert host["NetworkMode"] == "none", "Reader must not join a network"
    assert not host.get("Privileged"), "Reader must not be privileged"
    assert set(host.get("CapDrop") or []) == {"ALL"}
    assert not host.get("CapAdd") and not host.get("Binds") and not attrs.get("Mounts")
    assert not host.get("Devices") and not host.get("DeviceRequests")
    assert not host.get("PortBindings"), "Reader must not publish any host port"
    assert "no-new-privileges:true" in (host.get("SecurityOpt") or [])
    assert host["Memory"] == 1024**3 and host["MemorySwap"] == 1024**3
    assert host["NanoCpus"] == 1_000_000_000 and host["PidsLimit"] == 96
    assert host["Tmpfs"] == {"/tmp": "rw,noexec,nosuid,size=512m,mode=1777"}


def _remove_owned(client: Any, owned: list[str]) -> None:
    """Only explicit IDs created by this check are ever cleanup targets."""
    for container_id in owned:
        try:
            container = client.containers.get(container_id)
        except docker.errors.NotFound:
            continue
        gateway._remove_worker(container)


def _assert_owned_removed(client: Any, owned: list[str]) -> None:
    for container_id in owned:
        try:
            client.containers.get(container_id)
        except docker.errors.NotFound:
            continue
        raise AssertionError("An owned preview container remained after gateway cleanup")


def _fixtures(client: Any, image: str, owned: list[str]) -> dict[str, bytes]:
    container = client.containers.create(
        image=image, name="ai-dataseek-viz-fixtures-" + uuid.uuid4().hex,
        entrypoint=["/app/.venv/bin/python"], command=["-c", FIXTURE_CODE], working_dir="/app",
        user="65534:65534", network_mode="none", read_only=True,
        cap_drop=["ALL"], security_opt=["no-new-privileges:true"],
        mem_limit="1g", memswap_limit="1g", nano_cpus=1_000_000_000, pids_limit=96,
        tmpfs={"/tmp": "rw,noexec,nosuid,size=512m,mode=1777"},
        environment={"OPENBLAS_NUM_THREADS": "1", "OMP_NUM_THREADS": "1", "MKL_NUM_THREADS": "1", "PYTHONDONTWRITEBYTECODE": "1"},
        labels={"ai-dataseek.component": "visualization-fixture-check"},
    )
    owned.append(container.id)
    try:
        _inspect_isolation(container)
        container.start()
        status = container.wait(timeout=40)
        if status.get("StatusCode") != 0:
            raise AssertionError("Synthetic fixture producer failed in the constrained image")
        output = container.logs(stdout=True, stderr=False)
        assert len(output) <= 4 * 1024 * 1024, "Synthetic fixture output exceeded its budget"
        result = json.loads(output)
        assert result["uid"] == 65534 and isinstance(result["files"], dict)
        return {name: base64.b64decode(value, validate=True) for name, value in result["files"].items()}
    finally:
        container.remove(force=True)


@contextlib.contextmanager
def _observe_real_worker_containers(owned: list[str], verified: list[str]):
    """Transparent inspection of real SDK creates, with no fake I/O or result."""
    original_factory = gateway.docker.from_env

    class Containers:
        def __init__(self, real):
            self.real = real

        def create(self, **kwargs):
            container = self.real.create(**kwargs)
            owned.append(container.id)
            _inspect_isolation(container)
            assert container.attrs["Config"]["Entrypoint"] == ["/usr/bin/timeout"]
            assert container.attrs["Config"]["Cmd"] == ["--signal=KILL", "55s", "/app/.venv/bin/python", "-m", "app.services.extended_visualization_worker"]
            assert container.attrs["HostConfig"]["AutoRemove"] is True
            assert container.attrs["HostConfig"]["LogConfig"]["Type"] == "none"
            verified.append(container.id)
            return container

        def get(self, name):
            return self.real.get(name)

    class Client:
        def __init__(self, real):
            self.real = real
            self.containers = Containers(real.containers)

        def close(self):
            return self.real.close()

    def factory(*args, **kwargs):
        return Client(original_factory(*args, **kwargs))

    gateway.docker.from_env = factory
    try:
        yield
    finally:
        gateway.docker.from_env = original_factory


def _office_external_fixture() -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("_rels/.rels", '<Relationships><Relationship TargetMode="External" Target="file:///forbidden/external"/></Relationships>')
    return stream.getvalue()


def _blocked_gateway_child(image: str, connection: Any) -> None:
    """Run the real gateway with a deterministic test-only blocked Python body.

    Only the trusted test body replaces the reader module: PID 1, its fixed
    55-second limit, auto-removal, attach I/O and every isolation control are
    the production gateway's. No file/native-parser exploit is used to hang it.
    This child is killed with SIGKILL, so its finally cleanup cannot run.
    """
    original_factory = gateway.docker.from_env

    class Container:
        def __init__(self, real):
            self.real = real

        def __getattr__(self, name):
            return getattr(self.real, name)

        def start(self):
            self.real.start()
            connection.send({"id": self.real.id, "started": time.monotonic()})

    class Containers:
        def __init__(self, real):
            self.real = real

        def create(self, **kwargs):
            assert kwargs["entrypoint"] == ["/usr/bin/timeout"]
            assert kwargs["command"] == ["--signal=KILL", "55s", "/app/.venv/bin/python", "-m", "app.services.extended_visualization_worker"]
            assert kwargs["auto_remove"] is True
            kwargs["command"] = kwargs["command"][:3] + ["-c", "import sys,time; sys.stdin.buffer.read(); time.sleep(120)"]
            real = self.real.create(**kwargs)
            # Record the exact owned ID before inspecting or starting, so the
            # parent can clean it even if preparation fails.
            connection.send({"id": real.id})
            _inspect_isolation(real)
            assert real.attrs["HostConfig"]["AutoRemove"] is True
            return Container(real)

        def get(self, name):
            return self.real.get(name)

    class Client:
        def __init__(self, real):
            self.real = real
            self.containers = Containers(real.containers)

        def close(self):
            self.real.close()

    gateway.docker.from_env = lambda **kwargs: Client(original_factory(**kwargs))
    try:
        gateway.run_extended_visualization_worker(image, b"x,y\n1,2\n", reader="tabular", kind="table", format="csv",
                                                 options={}, truncated=False, cancelled=threading.Event())
    finally:
        gateway.docker.from_env = original_factory
        connection.close()


def _check_gateway_crash_cleanup(client: Any, image: str, owned: list[str]) -> dict:
    """Prove Docker reaps the worker after its separate caller is SIGKILLed."""
    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    process = context.Process(target=_blocked_gateway_child, args=(image, send))
    process.start()
    send.close()
    try:
        ready, startup_deadline = None, time.monotonic() + 15
        while time.monotonic() < startup_deadline:
            if receive.poll(.2):
                message = receive.recv()
                if message["id"] not in owned:
                    owned.append(message["id"])
                if "started" in message:
                    ready = message
                    break
            if not process.is_alive():
                raise AssertionError("Temporary crash-test gateway exited before startup")
        assert ready is not None, "Temporary crash-test gateway did not start"
        container = client.containers.get(ready["id"])
        _inspect_isolation(container)
        assert container.attrs["State"]["Running"] is True
        assert container.attrs["Config"]["Entrypoint"] == ["/usr/bin/timeout"]
        assert container.attrs["Config"]["Cmd"][:3] == ["--signal=KILL", "55s", "/app/.venv/bin/python"]
        assert container.attrs["HostConfig"]["AutoRemove"] is True
        process.kill()
        process.join(timeout=3)
        assert process.exitcode == -9, "Only the temporary gateway must be SIGKILLed"
        # No finally block from the dead gateway can remove this running worker.
        container.reload()
        assert container.attrs["State"]["Running"] is True
        deadline = ready["started"] + gateway.CONTAINER_TIMEOUT_SECONDS + 8
        while time.monotonic() < deadline:
            try:
                client.containers.get(ready["id"])
            except docker.errors.NotFound:
                elapsed = time.monotonic() - ready["started"]
                assert 50 <= elapsed <= 63, "Worker did not survive until its independent fixed deadline"
                return {"status": "passed", "gateway_exit_signal": "SIGKILL", "pid1": "/usr/bin/timeout",
                        "fixed_timeout_seconds": 55, "automatic_removal_seconds": round(elapsed, 3),
                        "blocked_body": "trusted synthetic Python sleep", "production_backend_untouched": True}
            time.sleep(.25)
        raise AssertionError("Worker survived its independent 55-second deadline after gateway crash")
    finally:
        if process.is_alive():
            process.kill()
            process.join(timeout=3)
        receive.close()
        process.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--image", help="Existing, already-built sandbox image (default: application settings)")
    parser.add_argument("--check-crash-cleanup", action="store_true",
                        help="Also SIGKILL a separate test gateway and wait for its real 55-second PID-1 cleanup")
    arguments = parser.parse_args()
    image = arguments.image or get_settings().sandbox_image
    if not image:
        raise SystemExit("Configure or explicitly select the existing sandbox image first")
    # Keep this factory client outside the transparent gateway observation.
    client = docker.from_env(timeout=5)
    owned, verified, checked, rejected = [], [], [], []
    started = time.monotonic()
    try:
        files = _fixtures(client, image, owned)
        fastq = b"".join(f"@read{i}\n".encode() + b"ACGT" * 25 + b"\n+\n" + b"I" * 100 + b"\n" for i in range(500))
        samples = [
            ("tabular", "table", "csv", b"x,y\n0,1.5\n1,2.5\n", {}),
            ("tabular", "heatmap", "npy", files["array.npy"], {"indices": [1]}),
            ("tabular", "heatmap", "npz", files["array.npz"], {}),
            ("tabular", "heatmap", "mat", files["array.mat"], {}),
            ("hdf5", "tree", "h5", files["sample.h5"], {}),
            ("hdf5", "heatmap", "nxs", files["sample.h5"], {"path": "/entry/signal", "indices": [1]}),
            ("hdf5", "series", "nc", files["NETCDF4.nc"], {"path": "/signal"}),
            ("excel", "table", "xlsx", files["book.xlsx"], {"sheet": "Measurements"}),
            ("rdkit", "image", "smi", b"CCO ethanol\n", {}),
            ("metpy", "image", "csv", b"p,t,td\n1000,20,16\n850,10,6\n700,0,-4\n500,-20,-25\n", METPY_OPTIONS),
            ("office", "pdf", "docx", files["document.docx"], {}),
            ("office", "pdf", "pptx", files["slides.pptx"], {}),
            ("root", "series", "root", files["histograms.root"], {"path": "/detector/counts"}),
            ("root", "heatmap", "root", files["histograms.root"], {"path": "/density"}),
            ("root", "series", "root", files["histograms.root"], {"path": "/curve"}),
            ("jcamp", "series", "jdx", JCAMP, {}),
            ("fastqc", "report", "fastq", fastq, {"confirm": True}),
        ]
        negatives = [
            ("hdf5-external-link", "hdf5", "series", "h5", files["sample.h5"], {"path": "/outside"}),
            ("classic-netcdf3-not-hdf5", "hdf5", "series", "nc", files["NETCDF3_CLASSIC.nc"], {}),
            ("office-external-reference", "office", "pdf", "docx", _office_external_fixture(), {}),
            ("root-executable-object-policy", "root", "series", "root", files["histograms.root"], {"path": "/description"}),
            ("jcamp-fid-policy", "jcamp", "series", "jdx", JCAMP.replace(b"NMR SPECTRUM", b"NMR FID"), {}),
            ("fastqc-explicit-confirmation", "fastqc", "report", "fastq", fastq, {"confirm": False}),
        ]
        with _observe_real_worker_containers(owned, verified):
            for reader, kind, fmt, data, options in samples:
                validate_options(reader, options)
                result = gateway.run_extended_visualization_worker(
                    image, data, reader=reader, kind=kind, format=fmt, options=options,
                    truncated=False, cancelled=threading.Event(),
                )
                assert result.get("ok") is True, f"{reader}/{kind}/{fmt}: " + result.get("error", "invalid envelope")
                payload = validate_payload(result["data"], reader, kind, MAX_OUTPUT)
                if reader == "excel":
                    assert payload["table"]["rows"][1] == [1, 20.5, None]
                    assert payload["table"]["formulas"][1] == [None, None, "=A2+B2"]
                if reader == "hdf5" and kind == "heatmap":
                    assert payload["array"]["values"] == list(range(12, 24))
                if reader == "root" and kind == "heatmap":
                    assert payload["array"]["shape"] == [3, 2]
                    assert payload["array"]["values"] == [0, 3, 1, 4, 2, 5]
                    assert payload["metadata"]["x_edges"] == [0, 1, 4]
                if reader == "jcamp":
                    assert payload["array"]["values"] == [4, 0, 3, 2, 2, 4, 1, 0]
                if reader == "fastqc":
                    basic = next(s for s in payload["sections"] if s["name"] == "Basic Statistics")
                    assert ["Total Sequences", "500"] in basic["rows"]
                    assert payload["metadata"]["engines"] == ["FastQC", "MultiQC"]
                    assert payload["table"]["rows"] and payload["metadata"]["complete_input"] is True
                checked.append({"reader": reader, "kind": kind, "format": fmt, "input_bytes": len(data),
                                "media_type": payload["media_type"], "array_shape": payload.get("array", {}).get("shape"),
                                "table_rows": len(payload.get("table", {}).get("rows", []))})
            for label, reader, kind, fmt, data, options in negatives:
                result = gateway.run_extended_visualization_worker(
                    image, data, reader=reader, kind=kind, format=fmt, options=options,
                    truncated=False, cancelled=threading.Event(),
                )
                assert result.get("ok") is False and isinstance(result.get("error"), str), label
                assert "Traceback" not in result["error"] and "/forbidden/" not in result["error"], label
                rejected.append(label)
            cancelled = threading.Event()
            timer = threading.Timer(.05, cancelled.set)
            timer.start()
            try:
                try:
                    gateway.run_extended_visualization_worker(image, files["document.docx"], reader="office", kind="pdf",
                                                             format="docx", options={}, truncated=False, cancelled=cancelled)
                except VisualizationWorkerError:
                    pass
                else:
                    raise AssertionError("Cancelled real worker must not return a preview")
            finally:
                timer.cancel()
        assert {item["reader"] for item in checked} == {"tabular", "hdf5", "excel", "rdkit", "metpy", "office", "root", "jcamp", "fastqc"}
        crash_cleanup = _check_gateway_crash_cleanup(client, image, owned) if arguments.check_crash_cleanup else {"status": "not requested"}
        _assert_owned_removed(client, owned)
        print(json.dumps({"status": "passed", "reader_count": 9, "previews": checked, "safe_rejections": rejected,
                          "actual_worker_isolation_verified": len(verified), "nonroot": "65534:65534", "memory_bytes": 1024**3,
                          "network": "none", "read_only_root": True, "host_mounts": 0, "published_ports": 0,
                          "cancellation": "passed", "owned_containers_remaining": 0,
                          "pid1_timeout_seconds": gateway.CONTAINER_TIMEOUT_SECONDS, "auto_remove": True,
                          "gateway_crash_cleanup": crash_cleanup,
                          "elapsed_seconds": round(time.monotonic() - started, 3),
                          "user_files_accessed": 0, "sessions_created": 0, "model_calls": 0}, ensure_ascii=False, indent=2))
    finally:
        _remove_owned(client, owned)
        client.close()


if __name__ == "__main__":
    main()
