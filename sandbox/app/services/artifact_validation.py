"""Bounded, read-only validation of output bytes, never model claims or code execution.

The receipt describes the exact bytes inspected. Callers must match its digest
to the bytes subsequently uploaded; path existence alone is not a receipt.
"""
from __future__ import annotations

import ast
import csv
import hashlib
import io
import json
import os
import re
import stat
import threading
import time
import warnings
import zipfile
from pathlib import Path
from typing import Any

from app.services.artifact_manifest import ARTIFACT_ROOT, MissingArtifact, _identity, _open_output

MAX_ITEMS = 32
MAX_FILE_BYTES = 32 * 1024 * 1024
MAX_BATCH_BYTES = 128 * 1024 * 1024
MAX_IMAGE_PIXELS = 16_000_000
MAX_IMAGE_FRAMES = 32
MAX_TABLE_ROWS = 100_000
MAX_TABLE_COLUMNS = 2_000
MAX_TABLE_CELLS = 200_000
MAX_MARKDOWN_TABLES = 128
MAX_ARCHIVE_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1_000
READ_CHUNK_BYTES = 1024 * 1024
BATCH_TIMEOUT_SECONDS = 30
KINDS = frozenset({"image", "table", "report", "code"})
IMAGE_FORMATS = {
    ".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP",
    ".gif": "GIF", ".bmp": "BMP", ".tif": "TIFF", ".tiff": "TIFF", ".avif": "AVIF",
}
VECTOR_IMAGE_SUFFIXES = frozenset({".svg"})
CODE_SUFFIXES = frozenset({".py", ".js", ".ts", ".r", ".sh", ".sql", ".ipynb"})
REPORT_SUFFIXES = frozenset({".txt", ".md", ".markdown", ".html", ".htm", ".json"})
DIAGNOSTIC_KEYS = frozenset({
    "row_number", "line_number", "column_number", "byte_offset", "expected_columns", "actual_columns",
    "row_count", "column_count", "size_bytes", "limit_bytes", "max_rows", "max_columns", "max_cells",
    "actual_cells", "width", "height", "frames", "max_pixels", "max_frames", "sheet_count", "max_sheets",
    "archive_entries", "max_archive_entries", "archive_bytes", "max_archive_bytes",
})


def _safe_diagnostics(value: Any) -> dict[str, int]:
    if not isinstance(value, dict):
        return {}
    return {key: value[key] for key in sorted(DIAGNOSTIC_KEYS) if key in value
            and type(value[key]) is int and 0 <= value[key] <= 2**63 - 1}


class _InvalidArtifact(Exception):
    def __init__(self, reason: str, diagnostics: dict | None = None):
        self.reason = reason
        self.diagnostics = _safe_diagnostics(diagnostics)


def _check_cancelled(cancelled: threading.Event, deadline: float) -> None:
    if cancelled.is_set() or time.monotonic() >= deadline:
        raise _InvalidArtifact("validation_deadline")


def _classify(path: str, expected_kind: str) -> str | None:
    suffix = Path(path).suffix.casefold()
    if suffix in IMAGE_FORMATS or suffix in VECTOR_IMAGE_SUFFIXES:
        return "image"
    if suffix in CODE_SUFFIXES:
        return "code"
    if suffix in {".csv", ".tsv", ".xlsx"}:
        return "table"
    if suffix == ".json" and expected_kind == "table":
        return "table"
    if suffix in {".md", ".markdown"} and expected_kind == "table":
        return "table"
    if suffix in REPORT_SUFFIXES:
        return "report"
    return None


def _read_stable_output(path: str, root: Path, cancelled: threading.Event,
                        deadline: float, byte_budget: list[int]) -> bytes:
    _check_cancelled(cancelled, deadline)
    with _open_output(path, root) as stream:
        before = os.fstat(stream.fileno())
        if not stat.S_ISREG(before.st_mode):
            raise _InvalidArtifact("not_regular_file")
        if before.st_size > MAX_FILE_BYTES:
            raise _InvalidArtifact("file_size_limit", {"size_bytes": before.st_size, "limit_bytes": MAX_FILE_BYTES})
        if before.st_size > byte_budget[0]:
            raise _InvalidArtifact("batch_size_limit", {"size_bytes": before.st_size, "limit_bytes": byte_budget[0]})
        if before.st_size == 0:
            raise _InvalidArtifact("empty_file")
        chunks: list[bytes] = []
        remaining = before.st_size + 1
        while remaining:
            _check_cancelled(cancelled, deadline)
            chunk = stream.read(min(READ_CHUNK_BYTES, remaining))
            if not chunk:
                break
            byte_budget[0] -= len(chunk)
            if byte_budget[0] < 0:
                raise _InvalidArtifact("batch_size_limit")
            chunks.append(chunk)
            remaining -= len(chunk)
        after = os.fstat(stream.fileno())
        # Re-traverse the root so a swapped parent cannot validate an old FD.
        try:
            with _open_output(path, root) as current_stream:
                current = os.fstat(current_stream.fileno())
        except (OSError, ValueError):
            # Once bytes have been read, disappearance or path replacement is
            # a changed artifact, never permission to recreate a missing one.
            raise _InvalidArtifact("changed_during_read") from None
        content = b"".join(chunks)
        if (_identity(before) != _identity(after) or _identity(after) != _identity(current)
                or len(content) != before.st_size):
            raise _InvalidArtifact("changed_during_read")
        return content


def _text(content: bytes) -> str:
    text = content.decode("utf-8-sig", errors="strict")
    if not text.strip() or "\x00" in text:
        raise _InvalidArtifact("empty_or_binary_text")
    return text


def _image_metadata(content: bytes, suffix: str, cancelled: threading.Event,
                    deadline: float) -> dict[str, Any]:
    from PIL import Image

    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(content)) as image:
            if image.format != IMAGE_FORMATS[suffix]:
                raise _InvalidArtifact("format_mismatch")
            image_format = image.format
            width, height = image.size
            if width < 1 or height < 1 or width * height > MAX_IMAGE_PIXELS:
                raise _InvalidArtifact("image_size_limit", {"width": width, "height": height, "max_pixels": MAX_IMAGE_PIXELS})
            image.verify()
        total_pixels = 0
        frames = 0
        with Image.open(io.BytesIO(content)) as image:
            # Verify and decode every permitted frame, not just a valid header.
            while True:
                _check_cancelled(cancelled, deadline)
                try:
                    image.seek(frames)
                except EOFError:
                    break
                frames += 1
                total_pixels += image.width * image.height
                if frames > MAX_IMAGE_FRAMES or total_pixels > MAX_IMAGE_PIXELS:
                    raise _InvalidArtifact("image_size_limit", {"frames": frames, "max_frames": MAX_IMAGE_FRAMES,
                                                               "max_pixels": MAX_IMAGE_PIXELS})
                image.load()
        return {"format": image_format, "width": width, "height": height, "frames": frames}


def _table_metadata(rows, cancelled: threading.Event, deadline: float) -> dict[str, int]:
    row_count = 0
    columns: int | None = None
    cells = 0
    for row_number, row in enumerate(rows, start=1):
        _check_cancelled(cancelled, deadline)
        if not row or all(value is None or value == "" for value in row):
            continue
        if columns is None:
            columns = len(row)
            if columns > MAX_TABLE_COLUMNS:
                raise _InvalidArtifact("table_size_limit", {"row_number": row_number, "column_count": columns,
                                                           "max_columns": MAX_TABLE_COLUMNS})
        elif len(row) != columns:
            raise _InvalidArtifact("inconsistent_table_width", {"row_number": row_number,
                "expected_columns": columns, "actual_columns": len(row)})
        row_count += 1
        cells += len(row)
        if row_count > MAX_TABLE_ROWS or cells > MAX_TABLE_CELLS:
            raise _InvalidArtifact("table_size_limit", {"row_count": row_count, "column_count": columns,
                "actual_cells": cells, "max_rows": MAX_TABLE_ROWS, "max_cells": MAX_TABLE_CELLS})
    if columns is None:
        raise _InvalidArtifact("empty_table")
    # Row count includes the header; no labels or cell values leave the sandbox.
    return {"row_count": row_count, "column_count": columns}


def _xlsx_metadata(content: bytes, cancelled: threading.Event, deadline: float) -> dict[str, Any]:
    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if (len(entries) > MAX_ARCHIVE_ENTRIES
                or sum(entry.file_size for entry in entries) > MAX_ARCHIVE_BYTES):
            raise _InvalidArtifact("archive_size_limit", {"archive_entries": len(entries),
                "max_archive_entries": MAX_ARCHIVE_ENTRIES, "archive_bytes": sum(entry.file_size for entry in entries),
                "max_archive_bytes": MAX_ARCHIVE_BYTES})
        for entry in entries:
            _check_cancelled(cancelled, deadline)
            if entry.flag_bits & 1:
                raise _InvalidArtifact("encrypted_workbook")
            if entry.filename.casefold().endswith((".xml", ".rels")):
                # UTF-16/32 XML is legal too: ignoring zero code units prevents
                # a DTD/entity declaration from hiding behind that encoding.
                xml = archive.read(entry).replace(b"\x00", b"").upper()
                if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                    raise _InvalidArtifact("unsafe_workbook_xml")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if not workbook.worksheets or len(workbook.worksheets) > 32:
            raise _InvalidArtifact("worksheet_count_limit", {"sheet_count": len(workbook.worksheets), "max_sheets": 32})
        sheets = []
        for sheet in workbook.worksheets:
            _check_cancelled(cancelled, deadline)
            if ((sheet.max_row or 0) > MAX_TABLE_ROWS or (sheet.max_column or 0) > MAX_TABLE_COLUMNS
                    or (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_TABLE_CELLS):
                raise _InvalidArtifact("table_size_limit", {"row_count": sheet.max_row or 0,
                    "column_count": sheet.max_column or 0, "max_rows": MAX_TABLE_ROWS,
                    "max_columns": MAX_TABLE_COLUMNS, "max_cells": MAX_TABLE_CELLS})
            try:
                sheets.append(_table_metadata(sheet.iter_rows(values_only=True), cancelled, deadline))
            except _InvalidArtifact as exc:
                if exc.reason != "empty_table":
                    raise
        if not sheets:
            raise _InvalidArtifact("empty_table")
        if sum(sheet["row_count"] * sheet["column_count"] for sheet in sheets) > MAX_TABLE_CELLS:
            raise _InvalidArtifact("table_size_limit")
        return {"format": "XLSX", "sheet_count": len(workbook.worksheets), "tables": sheets}
    finally:
        workbook.close()


def _strict_json(text: str) -> Any:
    def invalid_constant(_value):
        raise _InvalidArtifact("invalid_json_constant")

    def unique_object(pairs):
        payload = {}
        for key, value in pairs:
            if key in payload:
                raise _InvalidArtifact("duplicate_json_key")
            payload[key] = value
        return payload

    return json.loads(text, parse_constant=invalid_constant, object_pairs_hook=unique_object)


def _json_metadata(text: str, expected_kind: str) -> dict[str, Any]:
    payload = _strict_json(text)
    if expected_kind != "table":
        return {"format": "JSON", "json_type": type(payload).__name__}
    if not isinstance(payload, list) or not payload:
        raise _InvalidArtifact("invalid_json_table")
    if len(payload) > MAX_TABLE_ROWS:
        raise _InvalidArtifact("table_size_limit", {"row_count": len(payload), "max_rows": MAX_TABLE_ROWS})
    first = payload[0]
    if isinstance(first, dict):
        columns = set(first)
        valid = all(isinstance(row, dict) and set(row) == columns for row in payload)
        width = len(columns)
    elif isinstance(first, list):
        width = len(first)
        valid = all(isinstance(row, list) and len(row) == width for row in payload)
    else:
        raise _InvalidArtifact("invalid_json_table")
    if not valid or not width:
        raise _InvalidArtifact("invalid_json_table")
    if width > MAX_TABLE_COLUMNS or width * len(payload) > MAX_TABLE_CELLS:
        raise _InvalidArtifact("table_size_limit", {"row_count": len(payload), "column_count": width,
            "max_columns": MAX_TABLE_COLUMNS, "max_cells": MAX_TABLE_CELLS})
    return {"format": "JSON", "row_count": len(payload), "column_count": width}


def _markdown_row(line: str) -> list[str]:
    """Split GFM pipe cells, respecting escaped delimiters (also in code spans).

    This is a bounded table-structure check, not a Markdown renderer. Code
    fences and indented code are handled by the caller. No cell content leaves
    the sandbox. A line without a delimiter cannot establish a pipe table.
    """
    text = line.strip()
    cells, start, escaped = [], 0, False
    for index, char in enumerate(text):
        if char == "|" and not escaped:
            cells.append(text[start:index].strip())
            start = index + 1
        escaped = not escaped if char == "\\" else False
    if not cells:
        return []
    cells.append(text[start:].strip())
    if text.startswith("|"):
        cells.pop(0)
    if not cells[-1] and text.endswith("|"):
        cells.pop()
    return cells


def _markdown_metadata(text: str, cancelled: threading.Event, deadline: float) -> dict[str, Any]:
    """Recognize complete pipe tables without interpreting prose as a table.

    Ragged rows render with silently dropped/added cells in many Markdown
    clients. Preserve the report as readable text, but never advertise it as
    a validated table if any recognized table is inconsistent. Explicit table
    validation rejects it with a safe structural diagnostic below.
    """
    tables, failures = [], []
    previous = None
    active = None
    fence = None
    total_cells = 0
    for number, line in enumerate(text.splitlines(), 1):
        _check_cancelled(cancelled, deadline)
        match = re.match(r"^ {0,3}(`{3,}|~{3,})(.*)$", line)
        if fence:
            if match and match[1][0] == fence[0] and len(match[1]) >= fence[1] and not match[2].strip():
                fence = None
            previous, active = None, None
            continue
        if match:
            fence = (match[1][0], len(match[1]))
            previous, active = None, None
            continue
        if line.startswith(("    ", "\t")):
            previous, active = None, None
            continue
        cells = _markdown_row(line)
        delimiter = bool(cells and all(re.fullmatch(r":?-+:?", value) for value in cells))
        if delimiter and previous is not None and active is None:
            columns = len(previous[1])
            active = {"row_count": 1, "column_count": columns}
            tables.append(active)
            total_cells += columns
            if len(cells) != columns:
                failures.append({"row_number": number, "expected_columns": columns, "actual_columns": len(cells)})
        elif active is not None:
            if not cells:
                active = None
            else:
                active["row_count"] += 1
                total_cells += len(cells)
                if len(cells) != active["column_count"]:
                    failures.append({"row_number": number, "expected_columns": active["column_count"],
                                     "actual_columns": len(cells)})
        if (len(tables) > MAX_MARKDOWN_TABLES or total_cells > MAX_TABLE_CELLS
                or any(t["row_count"] > MAX_TABLE_ROWS or t["column_count"] > MAX_TABLE_COLUMNS
                       for t in tables[-1:])):
            raise _InvalidArtifact("table_size_limit")
        previous = (number, cells) if cells and not delimiter else None
    # A header and separator alone do not demonstrate populated table content.
    populated = [item for item in tables if item["row_count"] > 1]
    return {"format": "MARKDOWN", "validation_level": "nonempty_text",
            "table_validation": "inconsistent" if failures else "rectangular" if populated else "absent",
            "table_count": len(populated), "tables": populated,
            **({"table_diagnostics": failures[0]} if failures else {})}


def _notebook_metadata(text: str, cancelled: threading.Event, deadline: float) -> dict[str, Any]:
    """Validate nbformat 4 cell/output structure without repair or execution.

    This intentionally reports notebook_structure, not full JSON-schema or
    execution validation: arbitrary metadata and MIME payload semantics remain
    outside the check. Version 4.4 remains valid without cell IDs; 4.5+ requires
    valid, unique IDs. We never add IDs or erase outputs to make a file pass.
    """
    value = _strict_json(text)
    def require(condition):
        if not condition:
            raise _InvalidArtifact("invalid_notebook")
    def source(v):
        return isinstance(v, str) or isinstance(v, list) and all(isinstance(s, str) for s in v)
    def count(v):
        return v is None or type(v) is int and v >= 0
    require(isinstance(value, dict))
    require(type(value.get("nbformat")) is int and value["nbformat"] == 4
            and type(value.get("nbformat_minor")) is int and value["nbformat_minor"] >= 0
            and isinstance(value.get("metadata"), dict) and isinstance(value.get("cells"), list))
    require(len(value["cells"]) <= MAX_TABLE_CELLS)
    ids, code_cells, executed_cells, output_count = set(), 0, 0, 0
    for cell in value["cells"]:
        _check_cancelled(cancelled, deadline)
        require(isinstance(cell, dict) and isinstance(cell.get("cell_type"), str)
                and cell["cell_type"] in {"markdown", "raw", "code"}
                and isinstance(cell.get("metadata"), dict) and source(cell.get("source")))
        if value["nbformat_minor"] >= 5 or "id" in cell:
            identity = cell.get("id")
            require(isinstance(identity, str) and re.fullmatch(r"[A-Za-z0-9_-]{1,64}", identity)
                    and identity not in ids)
            ids.add(identity)
        if "attachments" in cell:
            require(cell["cell_type"] in {"markdown", "raw"} and isinstance(cell["attachments"], dict)
                    and all(isinstance(bundle, dict) for bundle in cell["attachments"].values()))
        if cell["cell_type"] != "code":
            require("outputs" not in cell and "execution_count" not in cell)
            continue
        code_cells += 1
        require("execution_count" in cell and count(cell["execution_count"]) and isinstance(cell.get("outputs"), list))
        executed_cells += cell["execution_count"] is not None
        output_count += len(cell["outputs"])
        require(output_count <= MAX_TABLE_CELLS)
        for output in cell["outputs"]:
            _check_cancelled(cancelled, deadline)
            require(isinstance(output, dict))
            kind = output.get("output_type")
            require(isinstance(kind, str))
            if kind == "stream":
                require(isinstance(output.get("name"), str) and output["name"] in {"stdout", "stderr"}
                        and source(output.get("text")))
            elif kind == "error":
                require(isinstance(output.get("ename"), str) and isinstance(output.get("evalue"), str)
                        and isinstance(output.get("traceback"), list)
                        and all(isinstance(s, str) for s in output["traceback"]))
            elif kind in {"display_data", "execute_result"}:
                require(isinstance(output.get("data"), dict) and isinstance(output.get("metadata"), dict))
                if kind == "execute_result":
                    require("execution_count" in output and count(output["execution_count"]))
            else:
                require(False)
    return {"format": "IPYNB", "validation_level": "notebook_structure", "nbformat": 4,
            "nbformat_minor": value["nbformat_minor"], "cell_count": len(value["cells"]),
            "code_cell_count": code_cells, "executed_cell_count": executed_cells, "output_count": output_count}


def _validate_content(content: bytes, suffix: str, kind: str, cancelled: threading.Event,
                      deadline: float) -> dict[str, Any]:
    _check_cancelled(cancelled, deadline)
    if suffix in VECTOR_IMAGE_SUFFIXES and kind == "image":
        from app.services.svg_validation import InvalidSvg, svg_metadata
        try:
            return svg_metadata(content, max_pixels=MAX_IMAGE_PIXELS,
                                check=lambda: _check_cancelled(cancelled, deadline), deadline=deadline)
        except InvalidSvg as error:
            raise _InvalidArtifact(error.reason, error.diagnostics) from None
    if kind == "image":
        return _image_metadata(content, suffix, cancelled, deadline)
    if suffix == ".xlsx":
        return _xlsx_metadata(content, cancelled, deadline)
    text = _text(content)
    if suffix == ".json":
        return _json_metadata(text, kind)
    if suffix in {".md", ".markdown"}:
        metadata = _markdown_metadata(text, cancelled, deadline)
        if kind == "table":
            if metadata["table_validation"] == "inconsistent":
                raise _InvalidArtifact("inconsistent_table_width", metadata["table_diagnostics"])
            if metadata["table_validation"] != "rectangular":
                raise _InvalidArtifact("empty_table")
        return metadata
    if kind == "table":
        rows = csv.reader(io.StringIO(text, newline=""), delimiter="\t" if suffix == ".tsv" else ",", strict=True)
        try:
            return {"format": suffix[1:].upper(), **_table_metadata(rows, cancelled, deadline)}
        except csv.Error:
            raise _InvalidArtifact("invalid_csv_syntax", {"line_number": rows.line_num}) from None
    if kind == "code":
        if suffix == ".py":
            ast.parse(text)  # Parse only: never compile or execute user code.
        elif suffix == ".ipynb":
            return _notebook_metadata(text, cancelled, deadline)
        return {"format": suffix[1:].upper(), "validation_level": "syntax" if suffix == ".py" else "nonempty_text"}
    return {"format": suffix[1:].upper(), "validation_level": "nonempty_text"}


def validate_artifacts(items: list[dict[str, str]], *, root: Path = ARTIFACT_ROOT,
                       cancelled: threading.Event | None = None, deadline: float | None = None) -> dict:
    """Return one receipt per requested file; never return file contents/errors."""
    if not items or len(items) > MAX_ITEMS:
        raise ValueError("invalid_item_count")
    cancelled = cancelled or threading.Event()
    deadline = time.monotonic() + BATCH_TIMEOUT_SECONDS if deadline is None else deadline
    files = []
    byte_budget = [MAX_BATCH_BYTES]
    for item in items:
        path, expected_kind = item["path"], item["kind"]
        kind = _classify(path, expected_kind)
        receipt = {"path": path, "expected_kind": expected_kind, "kind": kind, "valid": False,
                   "reason": "invalid_content", "sha256": None, "size": None, "metadata": {}, "diagnostics": {}}
        try:
            _check_cancelled(cancelled, deadline)
            if expected_kind not in KINDS:
                raise _InvalidArtifact("unsupported_kind")
            if kind is None:
                raise _InvalidArtifact("unsupported_format")
            if kind != expected_kind:
                raise _InvalidArtifact("kind_mismatch")
            try:
                content = _read_stable_output(path, root, cancelled, deadline, byte_budget)
            except MissingArtifact:
                raise _InvalidArtifact("missing_artifact") from None
            except (OSError, ValueError):
                # ELOOP/ENOTDIR and path validation failures are unsafe paths,
                # not malformed file contents eligible for local repair.
                # Keep this catch around reading only: image parsers may also
                # raise OSError when genuine output bytes are malformed.
                raise _InvalidArtifact("unavailable_or_unsafe_path") from None
            receipt["size"] = len(content)
            receipt["sha256"] = hashlib.sha256(content).hexdigest()
            metadata = _validate_content(content, Path(path).suffix.casefold(), kind, cancelled, deadline)
            _check_cancelled(cancelled, deadline)
            receipt.update(valid=True, reason="validated", metadata=metadata)
        except _InvalidArtifact as exc:
            receipt["reason"] = exc.reason
            receipt["diagnostics"] = exc.diagnostics
        except UnicodeDecodeError as exc:
            receipt.update(reason="invalid_text_encoding", diagnostics=_safe_diagnostics({"byte_offset": exc.start}))
        except json.JSONDecodeError as exc:
            receipt.update(reason="invalid_json_syntax", diagnostics=_safe_diagnostics({
                "line_number": exc.lineno, "column_number": exc.colno}))
        except SyntaxError as exc:
            receipt.update(reason="invalid_code_syntax", diagnostics=_safe_diagnostics({
                "line_number": exc.lineno, "column_number": exc.offset}))
        except (FileNotFoundError, PermissionError, IsADirectoryError):
            receipt["reason"] = "unavailable_or_unsafe_path"
        except ImportError:
            receipt["reason"] = "validator_unavailable"
        except Exception:
            # Parser/OS messages may contain paths, values or excerpts. The wire
            # response deliberately contains fixed reason codes only.
            receipt["reason"] = "invalid_content"
        files.append(receipt)
    return {"version": 1, "files": files}
