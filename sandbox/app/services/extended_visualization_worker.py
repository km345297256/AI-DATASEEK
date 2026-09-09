"""Bounded, one-shot v2 readers for trusted visualization adapters.

This module is run by the API's networkless, read-only worker container, never
inside the API process. The caller supplies bytes, an extension enum and typed
options, not a filesystem path, URL, command or Python expression. Optional
libraries fail closed. Third-party HTML and parser diagnostics are not returned.
"""
from __future__ import annotations

import base64
import csv
import io
import json
import math
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import zipfile
from datetime import date, datetime, time as datetime_time
from pathlib import Path
from xml.etree import ElementTree

from app.services.visualization_worker import _parser_output_to_stderr

MAX_INPUT_BYTES = 64 * 1024 * 1024
MAX_OUTPUT_BYTES = 8 * 1024 * 1024
MAX_MEDIA_BYTES = 5 * 1024 * 1024
MAX_ARRAY_VALUES = 16384
MAX_ROWS = 200
MAX_COLUMNS = 100
MAX_NODES = 256
MAX_DEPTH = 8
SUBPROCESS_SECONDS = 40
OFFICE_PAGE_LIMIT = 100
# The generated PDF is already size-bounded before parsing; cap its page tree
# as well. Older LibreOffice releases ignore CLI PageRange and may export more.
MAX_OFFICE_PDF_PAGES = 1000
OFFICE_PDF_SETTINGS = {
    "UseLosslessCompression": True,
    "ReduceImageResolution": False,
    "EmbedStandardFonts": True,
    "ExportBookmarks": False,
    "ExportFormFields": False,
}
FORMATS = {
    "tabular": {"csv", "tsv", "npy", "npz", "mat"},
    "hdf5": {"h5", "hdf5", "nxs", "nx", "nc4", "nc"},
    "excel": {"xlsx", "xls"},
    "rdkit": {"smi", "smiles", "mol", "sdf"},
    "metpy": {"csv", "tsv"},
    "office": {"docx", "doc", "pptx", "ppt", "odt", "odp"},
    "fastqc": {"fastq", "fq"},
    "root": {"root"}, "jcamp": {"jdx", "dx", "jcamp"},
}
KINDS = {
    "tabular": {"table", "series", "heatmap"},
    "hdf5": {"tree", "series", "heatmap"},
    "excel": {"table"}, "rdkit": {"image"}, "metpy": {"image"},
    "office": {"pdf"}, "fastqc": {"report"},
    "root": {"series", "heatmap"}, "jcamp": {"series"},
}
OPTIONS = {
    "tabular": {"variable", "row_offset", "column_offset", "x_column", "y_columns", "indices"},
    "hdf5": {"path", "indices"}, "excel": {"sheet", "row_offset", "column_offset"},
    "rdkit": {"molecule"},
    "metpy": {"pressure_column", "temperature_column", "dewpoint_column", "pressure_unit", "temperature_unit", "dewpoint_unit"},
    "office": set(), "fastqc": {"confirm"},
    "root": {"path"}, "jcamp": set(),
}


class PreviewError(ValueError):
    """Only fixed, non-sensitive messages may cross this boundary."""


def _label(value, limit=256):
    text = str(value or "")[:limit]
    if re.search(r"(?:/Users/|/home/|/tmp/|/private/|/var/|file:|https?://|[A-Za-z]:\\)", text, re.I):
        return "[redacted]"
    return "".join(c for c in text if ord(c) >= 32)


def _name(value):
    return (isinstance(value, str) and 0 < len(value) <= 128
            and _label(value, 128) == value and not any(c in value for c in "\\\x00"))


def _integer(value, maximum):
    return type(value) is int and 0 <= value <= maximum


def _options(reader, value):
    if not isinstance(value, dict) or set(value) - OPTIONS[reader]:
        raise PreviewError("不支持的可视化参数。")
    for key in ("variable", "sheet", "pressure_column", "temperature_column", "dewpoint_column"):
        if key in value and (not _name(value[key]) or "/" in value[key]):
            raise PreviewError("变量、工作表或列名称无效。")
    for key, maximum in (("row_offset", 100000), ("column_offset", 16383), ("x_column", 99), ("molecule", 99)):
        if key in value and not _integer(value[key], maximum):
            raise PreviewError("窗口或对象索引无效。")
    if "indices" in value and (not isinstance(value["indices"], list) or len(value["indices"]) > 6
                               or any(not _integer(v, 2**31 - 1) for v in value["indices"])):
        raise PreviewError("数组切片索引无效。")
    if "y_columns" in value and (not isinstance(value["y_columns"], list) or not 1 <= len(value["y_columns"]) <= 8
                                 or any(not _integer(v, 99) for v in value["y_columns"])
                                 or len(set(value["y_columns"])) != len(value["y_columns"])):
        raise PreviewError("曲线列选择无效。")
    if "path" in value and not _internal_path(value["path"]):
        raise PreviewError("HDF5 内部节点路径无效。")
    if reader == "metpy":
        if set(value) != OPTIONS[reader] or value["pressure_unit"] not in {"Pa", "hPa"} or any(
            value[k] not in {"degC", "K"} for k in ("temperature_unit", "dewpoint_unit")
        ):
            raise PreviewError("探空图必须明确指定气压、温度、露点列及其单位。")
        if len({value[k] for k in ("pressure_column", "temperature_column", "dewpoint_column")}) != 3:
            raise PreviewError("气压、温度和露点必须使用不同的数据列。")
    if reader == "fastqc" and (set(value) != {"confirm"} or value.get("confirm") is not True):
        raise PreviewError("完整质量分析必须由用户显式启动。")
    return value


def _internal_path(value):
    return (isinstance(value, str) and 0 < len(value) <= 512 and value.startswith("/")
            and _label(value, 512) == value and "\\" not in value
            and (value == "/" or all(_name(part) and part not in {".", ".."} for part in value[1:].split("/")))
            and value.count("/") <= MAX_DEPTH)


def _base(reader, kind):
    return {"contract_version": 2, "type": reader, "reader": reader, "kind": kind,
            "media_type": "application/json", "metadata": {}, "warnings": [], "sampled": False}


def _cell(value):
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else None
    return _label(value, 512)


def _numbers(values):
    import numpy as np
    array = np.ma.asarray(values, dtype=float).filled(float("nan"))
    if array.size > MAX_ARRAY_VALUES:
        raise PreviewError("数组超过交互式显示预算。")
    return [float(v) if math.isfinite(float(v)) else None for v in array.flat]


def _numeric_array(array):
    import numpy as np
    dtype = np.dtype(array.dtype)
    if dtype.kind not in "biuf" or not 1 <= len(array.shape) <= 8 or not all(array.shape):
        raise PreviewError("当前视图只支持非空、非复数的有界数值数组。")
    if math.prod(array.shape) * dtype.itemsize > MAX_INPUT_BYTES:
        raise PreviewError("数组解压后的大小超过安全预算。")


def _npy_preflight(stream, size):
    """Validate shape/dtype before numpy allocates an array from an archive."""
    import numpy as np
    version = np.lib.format.read_magic(stream)
    if version not in {(1, 0), (2, 0), (3, 0)}:
        raise PreviewError("NPY 协议版本不受支持。")
    shape, _, dtype = np.lib.format._read_array_header(stream, version, max_header_size=16384)
    if (dtype.kind not in "biuf" or not 1 <= len(shape) <= 8
            or any(type(v) is not int or v <= 0 for v in shape)
            or math.prod(shape) * dtype.itemsize > MAX_INPUT_BYTES
            or stream.tell() + math.prod(shape) * dtype.itemsize != size):
        raise PreviewError("NPY 数组类型、尺寸或数据长度不符合安全预算。")


def _array_result(reader, kind, array, options):
    import numpy as np
    _numeric_array(array)
    result = _base(reader, kind)
    leading = max(0, array.ndim - (1 if kind == "series" else 2))
    indices = options.get("indices", [])
    if len(indices) > leading or any(v >= array.shape[i] for i, v in enumerate(indices)):
        raise PreviewError("数组切片索引超出维度范围。")
    fixed = indices + [0] * (leading - len(indices))
    tail = array.shape[leading:]
    axis_limit = MAX_ARRAY_VALUES if len(tail) == 1 else 128
    steps = [max(1, math.ceil(n / axis_limit)) for n in tail]
    selection = tuple(fixed) + tuple(slice(None, None, step) for step in steps)
    selected = np.asarray(array[selection])
    result["array"] = {"shape": list(selected.shape), "dimensions": [f"dim_{i}" for i in range(leading, array.ndim)],
                       "values": _numbers(selected)}
    result["selected"] = {"indices": fixed}
    result["metadata"] = {"source_shape": list(array.shape), "strides": steps,
                          "axis_coordinates": "index", "selection": "strided sample; no aggregation"}
    result["sampled"] = bool(fixed) or any(step > 1 for step in steps)
    if result["sampled"]:
        result["warnings"].append("显示固定切片与等间隔抽样；未作全数组聚合，横轴为索引。")
    return result


def _zip_preflight(data, *, office=False):
    archive = zipfile.ZipFile(io.BytesIO(data))
    entries = archive.infolist()
    if not 1 <= len(entries) <= 4096 or sum(i.file_size for i in entries) > MAX_INPUT_BYTES:
        archive.close()
        raise PreviewError("压缩内容超过安全读取预算。")
    seen = set()
    for info in entries:
        name = info.filename
        if name in seen or name.startswith(("/", "\\")) or "\\" in name or ".." in name.split("/") or info.flag_bits & 1:
            archive.close()
            raise PreviewError("压缩文件包含不安全或加密的项目。")
        seen.add(name)
        if info.file_size > MAX_INPUT_BYTES or (info.compress_size and info.file_size / info.compress_size > 1000):
            archive.close()
            raise PreviewError("压缩项目超过解压预算。")
        lower = name.lower()
        if office and any(token in lower for token in ("vbaproject", "activex/", "embeddings/", "scripts/", "basic/", "externallinks/")):
            archive.close()
            raise PreviewError("预览拒绝包含宏、嵌入对象或外部链接的办公文件。")
        if office and (lower.endswith(".rels") or lower in {"content.xml", "styles.xml"}):
            if info.file_size > 8 * 1024 * 1024:
                archive.close()
                raise PreviewError("办公文档关系结构超过安全预算。")
            xml = archive.read(info)
            if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                archive.close()
                raise PreviewError("办公文档不允许外部 XML 实体。")
            root = ElementTree.fromstring(xml)
            for node in root.iter():
                if node.attrib.get("TargetMode", "").lower() == "external":
                    archive.close()
                    raise PreviewError("预览拒绝包含外部链接的办公文件。")
                if lower in {"content.xml", "styles.xml"} and any(
                    k.endswith("}href") and not (v.startswith("#") or re.fullmatch(r"(?:Pictures|ObjectReplacements)/[A-Za-z0-9_.-]+", v))
                    for k, v in node.attrib.items()
                ):
                    archive.close()
                    raise PreviewError("预览拒绝包含外部引用的办公文件。")
    return archive


def _csv_window(data, fmt, row_offset=0, column_offset=0):
    if b"\0" in data:
        raise PreviewError("表格不是有效的 UTF-8 文本。")
    stream = io.StringIO(data.decode("utf-8-sig"), newline="")
    csv.field_size_limit(65536)
    rows = csv.reader(stream, delimiter="\t" if fmt == "tsv" else ",")
    header = next(rows, None)
    if not header or len(header) > 16384 or any(len(v) > 65536 for v in header):
        raise PreviewError("表格缺少列名或列数超出限制。")
    selected = []
    total = 0
    more = False
    for row in rows:
        if len(row) > 16384:
            raise PreviewError("表格列数超出限制。")
        if total >= row_offset + MAX_ROWS:
            more = True
            break
        if total >= row_offset:
            selected.append([_cell(row[i]) if i < len(row) else None
                             for i in range(column_offset, min(len(header), column_offset + MAX_COLUMNS))])
        total += 1
    if column_offset >= len(header) or (row_offset and row_offset >= total):
        raise PreviewError("表格窗口超出数据范围。")
    return {"columns": [_label(v, 128) for v in header[column_offset:column_offset + MAX_COLUMNS]],
            "rows": selected, "row_offset": row_offset, "column_offset": column_offset,
            "total_rows": None if more else total, "total_columns": len(header)}, more


def tabular_preview(data, fmt, kind, options, path):
    import numpy as np
    variable = options.get("variable")
    names = []
    if fmt in {"csv", "tsv"}:
        if variable or "indices" in options:
            raise PreviewError("文本表格不支持数组变量和切片参数。")
        table, more = _csv_window(data, fmt, options.get("row_offset", 0), options.get("column_offset", 0))
        result = _base("tabular", kind)
        result["table"], result["sampled"] = table, more or table["column_offset"] > 0 or table["total_columns"] > MAX_COLUMNS
        if kind != "table":
            x_col, y_cols = options.get("x_column"), options.get("y_columns")
            width = len(table["columns"])
            if x_col is not None and x_col >= width or y_cols is not None and any(v >= width for v in y_cols):
                raise PreviewError("曲线列超出当前表格窗口。")
            columns = y_cols or list(range(min(8 if kind == "series" else 100, width)))
            values = []
            for row in table["rows"]:
                for column in columns:
                    try:
                        value = float(row[column]) if row[column] not in (None, "") else float("nan")
                    except (ValueError, TypeError):
                        raise PreviewError("选定的绘图列包含非数值内容，请明确选择数值列。") from None
                    values.append(value if math.isfinite(value) else None)
            if len(values) > MAX_ARRAY_VALUES:
                raise PreviewError("图形数据超过显示预算，请缩小列选择。")
            result["array"] = {"shape": [len(table["rows"]), len(columns)],
                               "dimensions": ["row", "column"], "values": values}
            result["selected"] = {"y_columns": columns, "x_column": x_col}
            result["metadata"]["axis_coordinates"] = "table window"
        return result
    if any(key in options for key in ("x_column", "y_columns")):
        raise PreviewError("数组视图不接受文本表格列参数。")
    if fmt == "npy":
        _npy_preflight(io.BytesIO(data), len(data))
        array = np.load(path, mmap_mode="r", allow_pickle=False, max_header_size=16384)
        names = ["array"]
        if variable not in {None, "array"}:
            raise PreviewError("数组变量不存在。")
    elif fmt == "npz":
        with _zip_preflight(data) as archive:
            members = archive.namelist()
            if len(members) > 128 or any(not n.endswith(".npy") or "/" in n or not _name(n[:-4]) for n in members):
                raise PreviewError("NPZ 只能包含至多 128 个命名数值数组。")
            names = [name[:-4] for name in members]
            variable = variable or names[0]
            if variable not in names:
                raise PreviewError("数组变量不存在。")
            member = variable + ".npy"
            with archive.open(member) as stream:
                _npy_preflight(stream, archive.getinfo(member).file_size)
            array = np.load(io.BytesIO(archive.read(member)), allow_pickle=False, max_header_size=16384)
    else:
        from scipy.io import loadmat, whosmat
        description = whosmat(path)
        if len(description) > 128:
            raise PreviewError("MAT 变量数量超过预览上限。")
        safe_classes = {"double", "single", "int8", "uint8", "int16", "uint16", "int32", "uint32", "int64", "uint64", "logical"}
        names = [name for name, shape, cls in description if _name(name) and "/" not in name and cls in safe_classes
                 and shape and math.prod(shape) * 8 <= MAX_INPUT_BYTES]
        variable = variable or next(iter(names), None)
        if variable not in names:
            raise PreviewError("MAT 未找到可预览的有界数值变量；v7.3 请使用 HDF5 插件。")
        array = loadmat(path, variable_names=[variable], verify_compressed_data_integrity=True)[variable]
    _numeric_array(array)
    if kind == "table":
        if array.ndim not in {1, 2} or "indices" in options:
            raise PreviewError("表格视图仅支持一维或二维数组。")
        matrix = array.reshape((-1, 1)) if array.ndim == 1 else array
        row, col = options.get("row_offset", 0), options.get("column_offset", 0)
        if row >= matrix.shape[0] or col >= matrix.shape[1]:
            raise PreviewError("表格窗口超出数组范围。")
        selected = matrix[row:row + MAX_ROWS, col:col + MAX_COLUMNS]
        result = _base("tabular", kind)
        result["table"] = {"columns": [f"column_{i}" for i in range(col, col + selected.shape[1])],
                           "rows": [[_cell(v.item()) for v in r] for r in selected],
                           "row_offset": row, "column_offset": col, "total_rows": matrix.shape[0], "total_columns": matrix.shape[1]}
        result["sampled"] = selected.shape != matrix.shape
    else:
        if any(k in options for k in ("row_offset", "column_offset")):
            raise PreviewError("数组绘图请使用切片参数，不支持表格窗口参数。")
        result = _array_result("tabular", kind, array, options)
    result["choices"] = {"variables": names}
    result.setdefault("selected", {})["variable"] = variable or "array"
    return result


def _hdf_attributes(node):
    import numpy as np
    result = {}
    for key in list(node.attrs)[:16]:
        if not _name(key) or "/" in key:
            continue
        aid = node.attrs.get_id(key)
        dtype = aid.dtype
        if dtype.kind not in "biufSU" or math.prod(aid.shape) * dtype.itemsize > 4096:
            continue
        value = np.asarray(node.attrs[key])
        if value.size > 16:
            continue
        cells = [_cell(v.decode("utf-8", "replace") if isinstance(v, bytes) else v.item() if hasattr(v, "item") else v) for v in value.flat]
        result[key] = cells[0] if value.ndim == 0 else cells
    return result


def _hdf_dataset_safe(node):
    if node.is_virtual or node.external:
        return False
    props = node.id.get_create_plist()
    if any(props.get_filter(i)[0] not in {1, 2, 3, 4, 5, 6, 32000} for i in range(props.get_nfilters())):
        return False
    return not node.chunks or math.prod(node.chunks) * node.dtype.itemsize <= MAX_INPUT_BYTES


def hdf5_preview(path, kind, options):
    import h5py
    result = _base("hdf5", kind)
    tree, datasets, seen = [], {}, set()
    with h5py.File(path, "r") as root:
        def walk(group, depth):
            if depth > MAX_DEPTH or len(tree) >= MAX_NODES:
                result["sampled"] = True
                return
            token = h5py.h5o.get_info(group.id).addr
            if token in seen:
                return
            seen.add(token)
            for key in group:
                if len(tree) >= MAX_NODES:
                    result["sampled"] = True
                    return
                internal = (group.name.rstrip("/") + "/" + key)
                if not _internal_path(internal):
                    continue
                link = group.get(key, getlink=True)
                if not isinstance(link, h5py.HardLink):
                    tree.append({"path": internal, "node_type": "blocked-link"})
                    continue
                node = group[key]
                if isinstance(node, h5py.Group):
                    tree.append({"path": internal, "node_type": "group", "attributes": _hdf_attributes(node)})
                    walk(node, depth + 1)
                elif isinstance(node, h5py.Dataset):
                    safe = _hdf_dataset_safe(node)
                    tree.append({"path": internal, "node_type": "dataset" if safe else "blocked-dataset",
                                 "shape": list(node.shape or ()), "dtype": _label(node.dtype, 64), "attributes": _hdf_attributes(node)})
                    if safe and node.dtype.kind in "biuf" and node.shape and 0 < len(node.shape) <= 8 and all(node.shape):
                        datasets[internal] = node
        walk(root, 1)
        if kind == "tree":
            if options:
                raise PreviewError("目录视图不接受数据切片参数。")
        else:
            selected = options.get("path") or next(iter(datasets), None)
            if selected not in datasets:
                raise PreviewError("未找到安全的数值节点；外部链接、虚拟数据集和自定义过滤器不可预览。")
            array_result = _array_result("hdf5", kind, datasets[selected], options)
            result.update(array_result)
            result["selected"]["path"] = selected
        result["tree"] = tree
        result["choices"] = {"variables": list(datasets)}
        if any(node["node_type"].startswith("blocked-") for node in tree):
            result["warnings"].append("已阻止软链接、外部链接、外部存储、虚拟数据集或不受信任的过滤器。")
        if len(tree) >= MAX_NODES:
            result["warnings"].append("目录仅列出前 256 个节点，深度最多 8 层。")
    return result


def _ole_preflight(data):
    import olefile
    if not olefile.isOleFile(io.BytesIO(data)):
        raise PreviewError("不是有效的旧版办公文件。")
    with olefile.OleFileIO(io.BytesIO(data)) as ole:
        entries = ole.listdir()
        if len(entries) > 4096:
            raise PreviewError("办公容器项目数量超出安全预算。")
        if any(any(token in part.lower() for token in ("vba", "macros", "objectpool", "encryptioninfo", "encryptedpackage"))
               for parts in entries for part in parts):
            raise PreviewError("预览拒绝包含宏、嵌入对象或加密内容的办公文件。")


def excel_preview(data, fmt, options, path):
    result = _base("excel", "table")
    row, col = options.get("row_offset", 0), options.get("column_offset", 0)
    if fmt == "xlsx":
        import openpyxl
        with _zip_preflight(data, office=True):
            pass
        with open(path, "rb") as stream, open(path, "rb") as formula_stream:
            book = openpyxl.load_workbook(stream, read_only=True, data_only=True, keep_links=False)
            formula_book = openpyxl.load_workbook(formula_stream, read_only=True, data_only=False, keep_links=False)
            try:
                names = book.sheetnames
                if len(names) > 256 or any(not _name(name) or "/" in name for name in names):
                    raise PreviewError("工作表数量或名称不符合预览限制。")
                selected = options.get("sheet") or names[0]
                if selected not in names:
                    raise PreviewError("所选工作表不存在。")
                sheet, formulas_sheet = book[selected], formula_book[selected]
                height, width = sheet.max_row or 0, sheet.max_column or 0
                if row >= height or col >= width or width > 16384 or height > 1048576:
                    raise PreviewError("工作表窗口超出有效范围。")
                window = dict(min_row=row + 1, max_row=min(height, row + MAX_ROWS), min_col=col + 1, max_col=min(width, col + MAX_COLUMNS))
                cells = [[_cell(v) for v in values] for values in sheet.iter_rows(**window, values_only=True)]
                formulas = [[_label(cell.value, 512) if cell.data_type == "f" else None for cell in values]
                            for values in formulas_sheet.iter_rows(**window)]
            finally:
                book.close()
                formula_book.close()
    else:
        import xlrd
        _ole_preflight(data)
        book = xlrd.open_workbook(file_contents=data, on_demand=True, formatting_info=False)
        try:
            names = book.sheet_names()
            if not 1 <= len(names) <= 256 or any(not _name(name) or "/" in name for name in names):
                raise PreviewError("工作表数量或名称不符合预览限制。")
            selected = options.get("sheet") or names[0]
            if selected not in names:
                raise PreviewError("所选工作表不存在。")
            sheet = book.sheet_by_name(selected)
            height, width = sheet.nrows, sheet.ncols
            if row >= height or col >= width:
                raise PreviewError("工作表窗口超出有效范围。")
            cells = []
            for r in range(row, min(height, row + MAX_ROWS)):
                current = []
                for c in range(col, min(width, col + MAX_COLUMNS)):
                    value = sheet.cell(r, c)
                    if value.ctype == xlrd.XL_CELL_DATE:
                        current.append(xlrd.xldate_as_datetime(value.value, book.datemode).isoformat())
                    elif value.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
                        current.append(None)
                    else:
                        current.append(_cell(value.value))
                cells.append(current)
            formulas = None
        finally:
            book.release_resources()
        result["warnings"].append("旧版 XLS 仅显示保存的计算值；此读取器不提供公式原文。")
    from openpyxl.utils import get_column_letter
    result["table"] = {"columns": [get_column_letter(c + 1) for c in range(col, min(width, col + MAX_COLUMNS))],
                       "rows": cells, "formulas": formulas, "row_offset": row, "column_offset": col,
                       "total_rows": height, "total_columns": width}
    result["choices"], result["selected"] = {"sheets": names}, {"sheet": selected}
    result["sampled"] = row > 0 or col > 0 or height > MAX_ROWS or width > MAX_COLUMNS
    result["metadata"]["formula_policy"] = "cached values only; no evaluation, macros or external links"
    result["warnings"].append("公式不重算；未保存的公式缓存显示为空，样式、图表和合并单元格不是原版式预览。")
    return result


def _media(result, data, media_type):
    if not 0 < len(data) <= MAX_MEDIA_BYTES:
        raise PreviewError("生成的预览文件超过显示预算。")
    result.update(media_type=media_type, data_base64=base64.b64encode(data).decode("ascii"))
    return result


def rdkit_preview(data, fmt, options):
    from rdkit import Chem
    from rdkit.Chem import Descriptors, rdMolDescriptors
    from rdkit.Chem.Draw import rdMolDraw2D
    if len(data) > 8 * 1024 * 1024:
        raise PreviewError("分子预览输入超过 8 MiB。")
    molecules, requested = [], options.get("molecule", 0)
    if fmt in {"smi", "smiles"}:
        for line in data.decode("utf-8-sig").splitlines():
            if line.strip() and not line.lstrip().startswith("#"):
                if len(line) > 65536:
                    raise PreviewError("SMILES 行长度超过限制。")
                molecules.append(Chem.MolFromSmiles(line.split()[0]))
                if len(molecules) >= 100:
                    break
    elif fmt == "mol":
        molecules = [Chem.MolFromMolBlock(data.decode("utf-8-sig"), strictParsing=True)]
    else:
        for molecule in Chem.ForwardSDMolSupplier(io.BytesIO(data), sanitize=True, removeHs=True, strictParsing=True):
            molecules.append(molecule)
            if len(molecules) >= 100:
                break
    if requested >= len(molecules) or molecules[requested] is None:
        raise PreviewError("所选分子无法解析或不在前 100 个对象中。")
    molecule = molecules[requested]
    if not 1 <= molecule.GetNumAtoms() <= 2000 or molecule.GetNumBonds() > 4000:
        raise PreviewError("分子规模超出二维结构预览预算。")
    drawer = rdMolDraw2D.MolDraw2DCairo(1000, 700)
    rdMolDraw2D.PrepareAndDrawMolecule(drawer, molecule)
    drawer.FinishDrawing()
    result = _base("rdkit", "image")
    result["metadata"] = {"formula": _label(rdMolDescriptors.CalcMolFormula(molecule)),
                          "molecular_weight": float(Descriptors.MolWt(molecule)), "atoms": molecule.GetNumAtoms(),
                          "bonds": molecule.GetNumBonds(), "molecules_listed": len(molecules)}
    result["selected"], result["sampled"] = {"molecule": requested}, len(molecules) >= 100
    return _media(result, drawer.GetDrawingText(), "image/png")


def metpy_preview(data, fmt, options):
    import numpy as np
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from metpy.plots import SkewT
    from metpy.units import units
    table, more = _csv_window(data, fmt)
    columns = [options[k] for k in ("pressure_column", "temperature_column", "dewpoint_column")]
    if any(table["columns"].count(name) != 1 for name in columns):
        raise PreviewError("探空数据缺少指定的唯一列名。")
    positions = [table["columns"].index(name) for name in columns]
    try:
        values = np.asarray([[float(row[i]) for i in positions] for row in table["rows"]])
    except (ValueError, TypeError):
        raise PreviewError("探空列必须包含完整数值，不能使用文本或缺失值。") from None
    if len(values) < 3 or not np.isfinite(values).all():
        raise PreviewError("探空图至少需要三个有效高度层。")
    pressure = units.Quantity(values[:, 0], options["pressure_unit"]).to("hPa")
    temperature = units.Quantity(values[:, 1], options["temperature_unit"]).to("degC")
    dewpoint = units.Quantity(values[:, 2], options["dewpoint_unit"]).to("degC")
    if (np.any(pressure.magnitude <= 0) or np.any(pressure.magnitude > 1200)
            or not (np.all(np.diff(pressure.magnitude) < 0) or np.all(np.diff(pressure.magnitude) > 0))
            or np.any(np.abs(temperature.magnitude) > 150) or np.any(np.abs(dewpoint.magnitude) > 150)):
        raise PreviewError("气压必须正值且严格单调；请检查变量映射和单位。")
    order = np.argsort(-pressure.magnitude)
    fig = plt.figure(figsize=(8, 8), dpi=100)
    try:
        skew = SkewT(fig, rotation=45)
        skew.plot(pressure[order], temperature[order], "r", label="Temperature")
        skew.plot(pressure[order], dewpoint[order], "g", label="Dew point")
        skew.ax.set_ylim(float(pressure.max().magnitude), max(50, float(pressure.min().magnitude)))
        skew.ax.set_xlim(-80, 50)
        skew.ax.legend()
        skew.ax.set_title("Skew-T: explicit columns and units")
        output = io.BytesIO()
        fig.savefig(output, format="png", bbox_inches="tight")
    finally:
        plt.close(fig)
    result = _base("metpy", "image")
    result["sampled"] = more
    result["selected"] = dict(options)
    result["metadata"] = {"levels": len(values), "diagram": "Skew-T", "diagnostic_calculations": False}
    result["warnings"].append("仅绘制最多前 200 个层次；不推断单位，不计算 CAPE/CIN，不替代气象业务诊断。")
    return _media(result, output.getvalue(), "image/png")


def _process(command, directory, timeout=SUBPROCESS_SECONDS):
    """Fixed argument arrays only; kill descendants on timeout, never retain logs."""
    env = {"PATH": os.environ.get("PATH", "/usr/local/bin:/usr/bin:/bin"), "LANG": "C.UTF-8", "LC_ALL": "C.UTF-8",
           "HOME": str(directory), "TMPDIR": str(directory), "MPLBACKEND": "Agg", "MPLCONFIGDIR": str(directory / "mpl"),
           "OPENBLAS_NUM_THREADS": "1", "OMP_NUM_THREADS": "1", "MULTIQC_DISABLE_ANONYMIZED_USAGE_REPORTING": "true"}
    process = subprocess.Popen(command, cwd=directory, env=env, stdin=subprocess.DEVNULL,
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
    try:
        status = process.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        process.wait(timeout=5)
        raise PreviewError("离线转换超过运行时间预算，请使用较小的文件。") from None
    if status:
        raise PreviewError("离线转换未成功完成；文件可能损坏、加密或使用不支持的功能。")


def _office_pdf_pages(pdf):
    """Bound the legacy converter's page range without rasterizing PDF content."""
    if len(pdf) > MAX_MEDIA_BYTES:
        raise PreviewError("高清办公 PDF 超过 5 MiB 预览预算；未降低清晰度，请拆分文档后重试。")
    from pypdf import PdfReader, PdfWriter

    try:
        reader = PdfReader(io.BytesIO(pdf), strict=True)
        count = reader.root_object["/Pages"]["/Count"]
        if not isinstance(count, int) or isinstance(count, bool) or not 1 <= count <= MAX_OFFICE_PDF_PAGES:
            raise PreviewError("办公 PDF 页数超过安全解析预算，请拆分文档后重试。")
        pages = len(reader.pages)
        if pages != count:
            raise PreviewError("办公 PDF 页数结构无效。")
        if pages <= OFFICE_PAGE_LIMIT:
            return pdf, pages, False
        output = io.BytesIO()
        writer = PdfWriter()
        # Copy page objects/embedded fonts/images verbatim. No rendering,
        # downsampling or JPEG round-trip occurs, including on LibreOffice 7.3.
        for page in reader.pages[:OFFICE_PAGE_LIMIT]:
            writer.add_page(page)
        writer.write(output)
        data = output.getvalue()
        if len(data) > MAX_MEDIA_BYTES:
            raise PreviewError("高清办公 PDF 超过 5 MiB 预览预算；未降低清晰度，请拆分文档后重试。")
        return data, OFFICE_PAGE_LIMIT, True
    except PreviewError:
        raise
    except Exception:
        raise PreviewError("办公转换 PDF 无法在安全预算内解析。") from None


def office_preview(data, fmt, directory):
    executable = shutil.which("soffice")
    if not executable:
        raise PreviewError("当前隔离镜像尚未安装 LibreOffice，无法转换此文件。")
    if fmt in {"docx", "pptx", "odt", "odp"}:
        with _zip_preflight(data, office=True) as archive:
            expected = {"docx": "word/document.xml", "pptx": "ppt/presentation.xml", "odt": "content.xml", "odp": "content.xml"}[fmt]
            if expected not in archive.namelist():
                raise PreviewError("办公文件内容与声明格式不一致。")
    else:
        _ole_preflight(data)
    source = directory / ("input." + fmt)
    if not source.exists():
        source.write_bytes(data)
    source.chmod(0o400)
    profile = directory / "profile"
    (profile / "user").mkdir(parents=True, mode=0o700)
    (profile / "user" / "registrymodifications.xcu").write_text(
        '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
        '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
        '<prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop>'
        '<prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop>'
        '<prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop></item>'
        '<item oor:path="/org.openoffice.Office.Writer/Content/Update">'
        '<prop oor:name="Link" oor:op="fuse"><value>2</value></prop>'
        '<prop oor:name="Field" oor:op="fuse"><value>false</value></prop>'
        '<prop oor:name="Chart" oor:op="fuse"><value>false</value></prop></item>'
        '<item oor:path="/org.openoffice.Office.Calc/Content/Update">'
        '<prop oor:name="Link" oor:op="fuse"><value>0</value></prop></item>'
        # CLI JSON filter options only work on LibreOffice 7.4+. The installed
        # 7.3 converter must also receive the same fixed settings in its private
        # profile. These are source-controlled, never file/user-supplied values.
        '<item oor:path="/org.openoffice.Office.Common/Filter/PDF/Export">'
        + ''.join('<prop oor:name="' + key + '" oor:op="fuse"><value>'
                  + str(value).lower() + '</value></prop>'
                  for key, value in OFFICE_PDF_SETTINGS.items())
        + '</item>'
        '</oor:items>', encoding="utf-8")
    output = directory / "output"
    output.mkdir(mode=0o700)
    filter_name = "impress_pdf_Export" if fmt in {"pptx", "ppt", "odp"} else "writer_pdf_Export"
    # Security is enforced by the private profile and relationship preflight,
    # in addition to the host's no-network/no-mount process isolation.
    pdf_options = json.dumps({"PageRange": {"type": "string", "value": f"1-{OFFICE_PAGE_LIMIT}"},
                              **{key: {"type": "boolean", "value": str(value).lower()}
                                 for key, value in OFFICE_PDF_SETTINGS.items()}}, separators=(",", ":"))
    _process([executable, "-env:UserInstallation=" + profile.as_uri(), "--headless", "--nologo", "--nodefault",
              "--norestore", "--unaccept=all", "--convert-to", "pdf:" + filter_name + ":" + pdf_options,
              "--outdir", str(output), str(source)], directory)
    target = output / "input.pdf"
    if not target.is_file() or target.is_symlink():
        raise PreviewError("办公转换未生成有效 PDF 文件。")
    if target.stat().st_size > MAX_MEDIA_BYTES:
        raise PreviewError("高清办公 PDF 超过 5 MiB 预览预算；未降低清晰度，请拆分文档后重试。")
    pdf = target.read_bytes()
    if not pdf.startswith(b"%PDF-"):
        raise PreviewError("办公转换结果不是有效 PDF。")
    pdf, pages, truncated = _office_pdf_pages(pdf)
    result = _base("office", "pdf")
    result["metadata"] = {"converter": "LibreOffice", "source_format": fmt, "page_limit": OFFICE_PAGE_LIMIT,
                          "page_count": pages, "image_compression": "lossless", "image_downsampling": False,
                          "standard_font_embedding": "requested", "macros": "disabled", "external_updates": "disabled", "editable": False}
    # A modern converter may already apply PageRange. At the exact limit the
    # original count is unknown, so conservatively signal a bounded preview.
    result["sampled"] = truncated or pages == OFFICE_PAGE_LIMIT
    result["warnings"].append("最多显示前 100 页静态版式；图像无损导出且不降采样，不播放动画、音视频，不执行宏，不更新外部链接。缺失字体会使用本地替代字体，与 Microsoft Office 排版可能不同。")
    return _media(result, pdf, "application/pdf")


def _fastqc_sections(text):
    sections, current = [], None
    for line in text.splitlines():
        if line.startswith(">>") and line != ">>END_MODULE":
            parts = line[2:].split("\t", 1)
            if len(sections) >= 32 or len(parts) != 2 or parts[1] not in {"pass", "warn", "fail"}:
                raise PreviewError("质量报告结构超出安全预算。")
            current = {"name": _label(parts[0], 128), "status": parts[1], "columns": [], "rows": []}
            sections.append(current)
        elif line == ">>END_MODULE":
            current = None
        elif current is not None and line.startswith("#"):
            current["columns"] = [_label(v, 128) for v in line[1:].split("\t")[:20]]
        elif current is not None and line:
            if len(current["rows"]) < 1000:
                current["rows"].append([_label(v, 256) for v in line.split("\t")[:20]])
    if not sections:
        raise PreviewError("质量分析未产生有效报告。")
    return sections


def fastqc_preview(data, directory):
    fastqc, multiqc = shutil.which("fastqc"), shutil.which("multiqc")
    if not fastqc or not multiqc:
        raise PreviewError("当前隔离镜像缺少 FastQC 或 MultiQC，无法执行完整质量分析。")
    # This is a complete bounded file, not the quick preview's first 2 MiB.
    source = directory / "sample.fastq"
    source.write_bytes(data)
    source.chmod(0o400)
    output = directory / "fastqc"
    output.mkdir(mode=0o700)
    started = time.monotonic()
    _process([fastqc, "--quiet", "--threads", "1", "--noextract", "--outdir", str(output), str(source)], directory, timeout=30)
    archive_path = output / "sample_fastqc.zip"
    if not archive_path.is_file() or archive_path.stat().st_size > 16 * 1024 * 1024:
        raise PreviewError("FastQC 报告不存在或超过预算。")
    with _zip_preflight(archive_path.read_bytes()) as archive:
        name = "sample_fastqc/fastqc_data.txt"
        if name not in archive.namelist() or archive.getinfo(name).file_size > 2 * 1024 * 1024:
            raise PreviewError("FastQC 结构化报告缺失或超出预算。")
        sections = _fastqc_sections(archive.read(name).decode("utf-8"))
    remaining = min(15, max(1, SUBPROCESS_SECONDS - (time.monotonic() - started)))
    aggregate = directory / "aggregate"
    aggregate.mkdir(mode=0o700)
    config = directory / "multiqc.yaml"
    config.write_text("disable_version_check: true\nno_ansi: true\n", encoding="utf-8")
    _process([multiqc, str(output), "--module", "fastqc", "--no-report", "--no-ansi", "--force",
              "--outdir", str(aggregate), "--config", str(config)], directory, timeout=remaining)
    summary_path = aggregate / "multiqc_data" / "multiqc_general_stats.txt"
    if not summary_path.is_file() or summary_path.stat().st_size > 128 * 1024:
        raise PreviewError("MultiQC 汇总报告缺失或超出预算。")
    summary, _ = _csv_window(summary_path.read_bytes(), "tsv")
    result = _base("fastqc", "report")
    result["sections"], result["table"] = sections, summary
    result["metadata"] = {"engines": ["FastQC", "MultiQC"], "complete_input": True, "input_bytes": len(data),
                          "samples": 1, "html_executed": False, "section_row_limit": 1000}
    result["warnings"].append("这是对当前单个完整有界文件的显式 QC；每个报告章节最多显示 1,000 行结构化结果，不运行第三方 HTML。")
    return result


def _xy_result(reader, kind, x, y):
    """Arrays carry full bounded data; optional table is only a small window."""
    import numpy as np
    x, y = np.asarray(x, dtype=float), np.asarray(y, dtype=float)
    if x.ndim != 1 or x.shape != y.shape or not 0 < len(x) <= MAX_ARRAY_VALUES // 2:
        raise PreviewError("坐标数据超过交互式曲线预算。")
    if not np.isfinite(x).all() or not np.isfinite(y).all():
        raise PreviewError("曲线坐标包含无效或非有限数值。")
    result = _base(reader, kind)
    result["array"] = {"shape": [len(x), 2], "dimensions": ["point", "coordinate"],
                       "values": _numbers(np.column_stack((x, y)))}
    if len(x) <= MAX_ROWS:
        result["table"] = {"columns": ["x", "y"], "rows": [[float(a), float(b)] for a, b in zip(x, y)],
                           "row_offset": 0, "column_offset": 0, "total_rows": len(x), "total_columns": 2}
    return result


def root_preview(path, kind, options):
    import numpy as np
    import uproot
    class FixedModelFile(uproot.ReadOnlyFile):
        @property
        def streamers(self):
            # Uproot can normally generate Python model classes from embedded
            # TStreamerInfo. This preview only accepts library-shipped models.
            raise PreviewError("ROOT 对象需要文件提供的动态流模型；当前插件仅允许内置固定模型。")
    classes = {f"TH{dimension}{storage}" for dimension in (1, 2) for storage in "CDFIS"}
    classes |= {"TGraph", "TGraphErrors", "TGraphAsymmErrors"}
    tree, allowed = [], {}
    with FixedModelFile(path, object_cache=None, array_cache=None, minimal_ttree_metadata=True,
                        use_threads=False, num_workers=1) as source:
        root = source.root_directory
        def walk(directory, prefix="", depth=1):
            if depth > MAX_DEPTH:
                return
            classnames = directory.classnames(recursive=False, cycle=False)
            if len(classnames) > MAX_NODES:
                raise PreviewError("ROOT 目录对象数量超过预览预算。")
            for name, cls in classnames.items():
                internal = prefix + "/" + name
                if not _internal_path(internal) or ";" in internal:
                    continue
                if len(tree) >= MAX_NODES:
                    raise PreviewError("ROOT 目录对象数量超过预览预算。")
                if cls in {"TDirectory", "TDirectoryFile"}:
                    tree.append({"path": internal, "node_type": "group", "dtype": cls})
                    walk(directory[name], internal, depth + 1)
                elif cls in classes:
                    tree.append({"path": internal, "node_type": "dataset", "dtype": cls})
                    allowed[internal] = cls
                else:
                    tree.append({"path": internal, "node_type": "blocked-object", "dtype": _label(cls, 64)})
        walk(root)
        choices = {p: cls for p, cls in allowed.items() if (cls.startswith("TH2")) == (kind == "heatmap")}
        selected = options.get("path") or next(iter(choices), None)
        if selected not in choices:
            raise PreviewError("请选择与视图匹配的 TH1、TH2 或 TGraph 数值对象；其他 ROOT 类不可执行或绘制。")
        key = root.key(selected[1:])
        if key.fObjlen > MAX_INPUT_BYTES:
            raise PreviewError("ROOT 对象解压大小超过安全预算。")
        obj, cls = root[selected[1:]], allowed[selected]
        if cls.startswith("TH1"):
            values, edges = np.asarray(obj.values(flow=False)), np.asarray(obj.axis(0).edges(flow=False))
            if values.ndim != 1 or not 0 < values.size <= 8192 or len(edges) != values.size + 1:
                raise PreviewError("ROOT 直方图分箱数量超过预览预算。")
            if not np.isfinite(edges).all() or not np.all(np.diff(edges) > 0):
                raise PreviewError("ROOT 分箱边界无效。")
            result = _xy_result("root", "series", (edges[:-1] + edges[1:]) / 2, values)
            result["array"] = {"shape": [values.size], "dimensions": ["x_bin"], "values": _numbers(values)}
            result["metadata"]["x_edges"] = _numbers(edges)
        elif cls.startswith("TH2"):
            values = np.asarray(obj.values(flow=False))
            x_edges, y_edges = np.asarray(obj.axis(0).edges(flow=False)), np.asarray(obj.axis(1).edges(flow=False))
            if (values.ndim != 2 or not 0 < values.size <= MAX_ARRAY_VALUES or max(values.shape) > 128
                    or not np.isfinite(values).all()
                    or len(x_edges) != values.shape[0] + 1 or len(y_edges) != values.shape[1] + 1
                    or not np.isfinite(x_edges).all() or not np.isfinite(y_edges).all()
                    or not np.all(np.diff(x_edges) > 0) or not np.all(np.diff(y_edges) > 0)):
                raise PreviewError("ROOT 二维分箱边界或规模不符合交互式预览预算。")
            result = _base("root", "heatmap")
            result["array"] = {"shape": list(values.T.shape), "dimensions": ["y_bin", "x_bin"], "values": _numbers(values.T)}
            result["metadata"].update(x_edges=_numbers(x_edges), y_edges=_numbers(y_edges))
        else:
            x, y = obj.values(axis="both")
            result = _xy_result("root", "series", x, y)
            if cls != "TGraph":
                result["warnings"].append("当前视图只显示 TGraph 坐标；误差条未纳入本期显示。")
        result["tree"], result["choices"], result["selected"] = tree, {"variables": list(allowed)}, {"path": selected}
        result["metadata"].update(root_class=cls, title=_label(getattr(obj, "title", "")), executable_objects=False)
        result["warnings"].append("只解析白名单数值对象；不执行 ROOT 宏、函数、画布命令。直方图不包含上溢与下溢分箱。")
        return result


_JCAMP_NUMBER = re.compile(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][-+]?\d+)?")


def _jcamp_numbers(text):
    remaining = _JCAMP_NUMBER.sub("", text)
    if remaining.strip(" \t,;") or len(text) > 65536:
        raise PreviewError("此 JCAMP 插件仅支持 AFFN/PAC 明文数值，不支持 SQZ/DIF/DUP 压缩。")
    values = [float(match.group()) for match in _JCAMP_NUMBER.finditer(text)]
    if not all(math.isfinite(value) for value in values):
        raise PreviewError("JCAMP 数值必须为有限值。")
    return values


def jcamp_preview(data):
    import numpy as np
    if len(data) > 8 * 1024 * 1024:
        raise PreviewError("JCAMP 输入超过 8 MiB 预算。")
    labels, lines, mode, ended = {}, [], None, False
    for raw in data.decode("utf-8-sig").splitlines():
        line = raw.split("$$", 1)[0].strip()
        if not line:
            continue
        if ended:
            raise PreviewError("当前插件只支持单个 JCAMP 谱块。")
        if line.startswith("##"):
            if "=" not in line:
                raise PreviewError("JCAMP 标签格式无效。")
            name, value = line[2:].split("=", 1)
            name, value = re.sub(r"[ _-]", "", name.upper()), value.strip()
            if len(name) > 64 or len(value) > 4096 or name in labels:
                raise PreviewError("JCAMP 标签重复或超出安全预算。")
            labels[name] = value
            if len(labels) > 256:
                raise PreviewError("JCAMP 标签数量超出安全预算。")
            if name == "END":
                ended = True
            elif name in {"XYDATA", "XYPOINTS"}:
                if mode is not None or value.replace(" ", "").upper() not in {"(X++(Y..Y))", "(XY..XY)"}:
                    raise PreviewError("仅支持一维明文 XYDATA 或 XYPOINTS 谱。")
                mode = name
                if (name == "XYDATA") != (value.replace(" ", "").upper() == "(X++(Y..Y))"):
                    raise PreviewError("JCAMP 数据布局声明不一致。")
            elif mode:
                raise PreviewError("不支持在数据段之后追加多谱或结构标签。")
        elif mode:
            lines.append(line)
            if len(lines) > 8192:
                raise PreviewError("JCAMP 谱点数量超出预览预算。")
        else:
            raise PreviewError("JCAMP 标签不支持跨行或未声明的数据。")
    if not ended or mode is None or labels.get("DATATYPE", "").upper().strip() != "NMR SPECTRUM":
        raise PreviewError("当前插件只支持已经处理的一维 NMR SPECTRUM，不接收 FID、NTUPLES 或二维数据。")
    if any(name in labels for name in ("NTUPLES", "DATATABLE", "NUMDIM", "VARNAME", "SYMBOL")):
        raise PreviewError("当前插件不支持多维、复数或 NTUPLES 数据。")
    if labels.get("XUNITS", "").upper().strip() != "PPM":
        raise PreviewError("NMR 预览必须明确使用 PPM 横轴；不自动换算频率或时间。")
    nucleus = labels.get(".OBSERVENUCLEUS", "").strip("<> ")
    if not re.fullmatch(r"\d{1,3}[A-Z][a-z]?", nucleus):
        raise PreviewError("NMR 谱必须提供明确的观测核，如 1H 或 13C。")
    def number(name, default=None):
        value = labels.get(name)
        if value is None:
            if default is None:
                raise PreviewError("JCAMP 缺少必要的点数或坐标标签。")
            return default
        if not _JCAMP_NUMBER.fullmatch(value):
            raise PreviewError("JCAMP 坐标标签不是单个明文数值。")
        parsed = float(value)
        if not math.isfinite(parsed):
            raise PreviewError("JCAMP 坐标标签包含无效数值。")
        return parsed
    count = number("NPOINTS")
    if count != int(count) or not 2 <= count <= MAX_ARRAY_VALUES // 2:
        raise PreviewError("JCAMP 谱点数必须介于 2 和 8192。")
    count, x_factor, y_factor = int(count), number("XFACTOR", 1.0), number("YFACTOR", 1.0)
    if x_factor == 0 or y_factor == 0:
        raise PreviewError("JCAMP 缩放系数不能为零。")
    xs, ys = [], []
    if mode == "XYDATA":
        first, last = number("FIRSTX"), number("LASTX")
        delta = (last - first) / (count - 1)
        if delta == 0:
            raise PreviewError("JCAMP 横轴坐标必须严格单调。")
        for line in lines:
            values = _jcamp_numbers(line)
            if len(values) < 2 or not math.isclose(values[0] * x_factor, first + len(ys) * delta, rel_tol=1e-5, abs_tol=abs(delta) * .01):
                raise PreviewError("JCAMP 行首坐标与点数、步长不一致。")
            ys.extend(value * y_factor for value in values[1:])
            if len(ys) > count:
                raise PreviewError("JCAMP 数据点超过声明数量。")
        xs = [first + i * delta for i in range(len(ys))]
    else:
        for line in lines:
            values = _jcamp_numbers(line)
            if len(values) % 2:
                raise PreviewError("JCAMP XYPOINTS 必须包含成对坐标。")
            xs.extend(value * x_factor for value in values[::2])
            ys.extend(value * y_factor for value in values[1::2])
            if len(ys) > count:
                raise PreviewError("JCAMP 数据点超过声明数量。")
    if len(ys) != count or not (np.all(np.diff(xs) > 0) or np.all(np.diff(xs) < 0)):
        raise PreviewError("JCAMP 点数不匹配或横轴不是严格单调。")
    result = _xy_result("jcamp", "series", xs, ys)
    result["metadata"] = {"is_fid": False, "x_unit": "PPM", "nucleus": nucleus,
                          "title": _label(labels.get("TITLE", "")), "source_points": count,
                          "encoding": "AFFN/PAC", "processed": True}
    if ".OBSERVEFREQUENCY" in labels:
        frequency = number(".OBSERVEFREQUENCY")
        if frequency <= 0:
            raise PreviewError("观测频率必须为正数。")
        result["metadata"]["frequency"] = frequency
    result["warnings"].append("仅读取已处理的一维明文 NMR 谱；不做傅里叶变换、相位/基线校正或频率单位推断。")
    return result


def preview_bytes(data, reader, kind, options=None, *, format=None, truncated=False):
    if reader not in FORMATS or kind not in KINDS[reader] or format not in FORMATS[reader]:
        raise PreviewError("未注册的扩展读取器、视图或格式。")
    options = _options(reader, {} if options is None else options)
    if type(truncated) is not bool or truncated:
        raise PreviewError("此扩展读取器需要完整的有界文件。")
    if not isinstance(data, bytes) or not 0 < len(data) <= MAX_INPUT_BYTES:
        raise PreviewError("文件为空或超出 64 MiB 读取预算。")
    # Do not allow HDF5 to discover executable filter plugins from the runtime.
    os.environ["HDF5_PLUGIN_PRELOAD"] = "::"
    os.environ["HDF5_PLUGIN_PATH"] = ""
    with tempfile.TemporaryDirectory(prefix="visualization-v2-") as temporary:
        directory = Path(temporary)
        directory.chmod(0o700)
        path = directory / ("input." + format)
        path.write_bytes(data)
        path.chmod(0o400)
        if reader == "tabular":
            return tabular_preview(data, format, kind, options, path)
        if reader == "hdf5":
            return hdf5_preview(path, kind, options)
        if reader == "excel":
            return excel_preview(data, format, options, path)
        if reader == "rdkit":
            return rdkit_preview(data, format, options)
        if reader == "metpy":
            return metpy_preview(data, format, options)
        if reader == "office":
            return office_preview(data, format, directory)
        if reader == "root":
            return root_preview(path, kind, options)
        if reader == "jcamp":
            return jcamp_preview(data)
        return fastqc_preview(data, directory)


def _encoded_result(stream=None):
    try:
        stream = sys.stdin.buffer if stream is None else stream
        line = stream.readline(8193)
        if len(line) > 8192 or not line.endswith(b"\n"):
            raise PreviewError("无效的预览协议头。")
        header = json.loads(line)
        if (not isinstance(header, dict)
                or set(header) - {"contract_version", "size", "reader", "kind", "format", "options", "truncated"}
                or type(header.get("contract_version")) is not int or header["contract_version"] != 2
                or type(header.get("size")) is not int or not 0 < header["size"] <= MAX_INPUT_BYTES):
            raise PreviewError("无效的预览协议版本或大小。")
        data = stream.read(header["size"])
        if len(data) != header["size"] or stream.read(1):
            raise PreviewError("预览数据长度与协议声明不一致。")
        payload = preview_bytes(data, header.get("reader"), header.get("kind"), header.get("options"),
                                format=header.get("format"), truncated=header.get("truncated", False))
        encoded = json.dumps({"ok": True, "data": payload}, ensure_ascii=False, allow_nan=False, separators=(",", ":")).encode()
        if len(encoded) > MAX_OUTPUT_BYTES:
            raise PreviewError("可视化结果超过 8 MiB 显示预算。")
    except PreviewError as error:
        encoded = json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False).encode()
    except ImportError:
        encoded = json.dumps({"ok": False, "error": "隔离镜像缺少所需读取依赖，请更新对应插件运行环境。"}, ensure_ascii=False).encode()
    except Exception:
        encoded = json.dumps({"ok": False, "error": "无法安全解析此文件；请检查文件格式或使用领域分析工具。"}, ensure_ascii=False).encode()
    return encoded


def main():
    with _parser_output_to_stderr():
        encoded = _encoded_result()
    sys.stdout.buffer.write(encoded)
    sys.stdout.buffer.flush()


if __name__ == "__main__":
    main()
